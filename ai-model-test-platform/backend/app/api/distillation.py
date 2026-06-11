from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from typing import List
import csv
import io
import json

from app.database import get_db
from app.models.distilled_user import DistilledUser
from app.models.chat_record import ChatRecord
from app.schemas.distilled_user import DistilledUserCreate, DistilledUserResponse, DistilledUserDetail

router = APIRouter(prefix="/api/distillation", tags=["distillation"])


def parse_csv_file(file_content: bytes):
    """解析CSV文件"""
    content = file_content.decode('utf-8')
    reader = csv.DictReader(io.StringIO(content))
    rows = list(reader)
    return rows


def parse_excel_file(file_content: bytes):
    """解析Excel文件（简化实现，使用openpyxl）"""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(file_content))
        ws = wb.active
        
        # 获取表头
        headers = [cell.value for cell in ws[1]]
        
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            row_dict = {}
            for i, value in enumerate(row):
                if i < len(headers):
                    row_dict[headers[i]] = value
            rows.append(row_dict)
        
        return rows
    except ImportError:
        raise HTTPException(status_code=500, detail="缺少openpyxl依赖，无法解析Excel文件")


@router.get("/users", response_model=List[DistilledUserResponse])
def list_users(skip: int = 0, limit: int = 100, db: Session = Depends(get_db)):
    """获取蒸馏用户列表"""
    users = db.query(DistilledUser).filter(DistilledUser.status == 1).offset(skip).limit(limit).all()
    return users


@router.post("/users", response_model=DistilledUserResponse)
def create_user(user: DistilledUserCreate, db: Session = Depends(get_db)):
    """创建蒸馏用户（手动输入）"""
    db_user = DistilledUser(name=user.name)
    db.add(db_user)
    db.commit()
    db.refresh(db_user)
    
    # 添加聊天记录
    if user.chat_records:
        for record in user.chat_records:
            db_record = ChatRecord(
                user_id=db_user.id,
                role=record.role,
                content=record.content,
                timestamp=record.timestamp
            )
            db.add(db_record)
        
        db_user.message_count = len(user.chat_records)
        db.commit()
    
    return db_user


@router.post("/users/upload")
def upload_chat_records(
    file: UploadFile = File(...),
    user_name: str = None,
    db: Session = Depends(get_db)
):
    """上传聊天记录文件（Excel/CSV）"""
    
    if not file.filename.endswith(('.xlsx', '.xls', '.csv')):
        raise HTTPException(status_code=400, detail="只支持Excel或CSV文件")
    
    try:
        content = file.file.read()
        
        # 读取文件
        if file.filename.endswith('.csv'):
            rows = parse_csv_file(content)
        else:
            rows = parse_excel_file(content)
        
        # 检查必要列
        if not rows:
            raise HTTPException(status_code=400, detail="文件为空")
        
        required_columns = ['role', 'content']
        if not all(col in rows[0] for col in required_columns):
            raise HTTPException(status_code=400, detail=f"文件必须包含列: {required_columns}")
        
        # 创建用户
        name = user_name or file.filename.rsplit('.', 1)[0]
        db_user = DistilledUser(
            name=name,
            source_file=file.filename,
            message_count=len(rows)
        )
        db.add(db_user)
        db.commit()
        db.refresh(db_user)
        
        # 添加聊天记录
        for row in rows:
            db_record = ChatRecord(
                user_id=db_user.id,
                role=row.get('role', 'user'),
                content=str(row.get('content', '')),
                timestamp=row.get('timestamp')
            )
            db.add(db_record)
        
        db.commit()
        
        return {
            "id": db_user.id,
            "name": db_user.name,
            "message_count": db_user.message_count,
            "message": "上传成功"
        }
        
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"文件解析失败: {str(e)}")


@router.get("/users/{user_id}", response_model=DistilledUserDetail)
def get_user(user_id: int, db: Session = Depends(get_db)):
    """获取用户详情"""
    user = db.query(DistilledUser).filter(DistilledUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    
    # 获取聊天记录
    records = db.query(ChatRecord).filter(ChatRecord.user_id == user_id).all()
    user.chat_records = [
        {"role": r.role, "content": r.content, "timestamp": r.timestamp}
        for r in records
    ]
    
    return user


@router.delete("/users/{user_id}")
def delete_user(user_id: int, db: Session = Depends(get_db)):
    """删除用户（软删除）"""
    user = db.query(DistilledUser).filter(DistilledUser.id == user_id).first()
    if not user:
        raise HTTPException(status_code=404, detail="用户不存在")
    
    user.status = 0
    db.commit()
    
    return {"message": "删除成功"}
